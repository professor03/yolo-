#!/usr/bin/env python3
"""
匯出功能 API 路由
提供 CSV 和 Excel 格式的數據匯出
"""

import csv
import io
from datetime import datetime, timedelta
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse

try:
    import openpyxl
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from src.auth.security import get_security_manager
from src.api.data_store import data_store
from src.database.models import get_database_manager, SystemMetrics
from src.database.dependencies import get_db_session
from sqlalchemy.orm import Session

# 創建路由器
router = APIRouter(prefix="/export", tags=["export"])

from fastapi import Header

def get_current_user_with_export_permission(authorization: str = Header(None)):
    """獲取有匯出權限的當前用戶"""
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="未提供認證令牌")
    
    token = authorization.split(" ")[1]
    security_manager = get_security_manager()
    payload = security_manager.verify_token(token, "access")
    
    if not payload:
        raise HTTPException(status_code=401, detail="無效的認證令牌")
    
    # 檢查匯出權限
    if not security_manager.check_permission(payload.get("role"), "data_export"):
        raise HTTPException(status_code=403, detail="需要數據匯出權限")
    
    return payload

def get_metrics_data(
    hours: int = 24,
    db: Session = None
) -> List[Dict]:
    """獲取指標數據"""
    if db is None:
        db_manager = get_database_manager()
        db = db_manager.get_session()
        should_close = True
    else:
        should_close = False
    
    try:
        # 計算時間範圍
        end_time = datetime.utcnow()
        start_time = end_time - timedelta(hours=hours)
        
        # 從資料庫查詢
        db_metrics = (
            db.query(SystemMetrics)
            .filter(SystemMetrics.timestamp >= start_time)
            .filter(SystemMetrics.timestamp <= end_time)
            .order_by(SystemMetrics.timestamp.desc())
            .limit(10000)  # 限制最大記錄數
            .all()
        )
        
        # 轉換為字典格式
        metrics_data = []
        for metric in db_metrics:
            metrics_data.append({
                "timestamp": metric.timestamp.isoformat() if metric.timestamp else "",
                "stream_id": metric.source_id or "",
                "people_count": metric.people_count or 0,
                "fps_processing": metric.fps_processing or 0.0,
                "fps_source": metric.fps_source or 0.0,
                "fps_inference": metric.fps_inference or 0.0,
                "latency_ms": metric.latency_ms or 0.0,
                "cpu_percent": metric.cpu_percent or 0.0,
                "memory_mb": metric.memory_mb or 0.0,
                "gpu_utilization": metric.gpu_utilization or 0.0,
                "gpu_memory_gb": metric.gpu_memory_gb or 0.0,
                "total_detections": metric.total_detections or 0,
                "queue_size": metric.queue_size or 0,
                "dropped_frames": metric.dropped_frames or 0,
                "session_id": metric.session_id or ""
            })
        
        # 如果資料庫沒有數據，嘗試從內存數據存儲獲取
        if not metrics_data:
            # 從 data_store 獲取歷史數據
            history = data_store.get_metrics_history()
            for entry in history[-1000:]:  # 最近1000條記錄
                metrics_data.append({
                    "timestamp": entry.get("timestamp", ""),
                    "stream_id": entry.get("stream_id", ""),
                    "people_count": entry.get("people_count", 0),
                    "fps_processing": entry.get("fps_processing", 0.0),
                    "fps_source": entry.get("fps_source", 0.0),
                    "fps_inference": entry.get("fps_inference", 0.0),
                    "latency_ms": entry.get("latency_ms", 0.0),
                    "cpu_percent": entry.get("cpu_percent", 0.0),
                    "memory_mb": entry.get("memory_mb", 0.0),
                    "gpu_utilization": 0.0,
                    "gpu_memory_gb": 0.0,
                    "total_detections": entry.get("total_detections", 0),
                    "queue_size": 0,
                    "dropped_frames": 0,
                    "session_id": ""
                })
        
        return metrics_data
    
    finally:
        if should_close:
            db.close()

@router.get("/metrics.csv")
async def export_metrics_csv(
    hours: int = Query(24, ge=1, le=168, description="匯出最近N小時的數據"),
    current_user: Dict = Depends(get_current_user_with_export_permission)
):
    """匯出指標數據為 CSV 格式"""
    
    # 獲取數據
    metrics_data = get_metrics_data(hours)
    
    if not metrics_data:
        # 如果沒有數據，返回空的 CSV
        metrics_data = [{
            "timestamp": datetime.utcnow().isoformat(),
            "stream_id": "no_data",
            "people_count": 0,
            "fps_processing": 0.0,
            "fps_source": 0.0,
            "fps_inference": 0.0,
            "latency_ms": 0.0,
            "cpu_percent": 0.0,
            "memory_mb": 0.0,
            "gpu_utilization": 0.0,
            "gpu_memory_gb": 0.0,
            "total_detections": 0,
            "queue_size": 0,
            "dropped_frames": 0,
            "session_id": ""
        }]
    
    # 創建 CSV 內容
    output = io.StringIO()
    fieldnames = [
        "timestamp", "stream_id", "people_count", "fps_processing", 
        "fps_source", "fps_inference", "latency_ms", "cpu_percent", 
        "memory_mb", "gpu_utilization", "gpu_memory_gb", "total_detections",
        "queue_size", "dropped_frames", "session_id"
    ]
    
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(metrics_data)
    
    # 準備響應
    csv_content = output.getvalue()
    output.close()
    
    filename = f"yolo_metrics_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
    
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )

@router.get("/metrics.xlsx")
async def export_metrics_xlsx(
    hours: int = Query(24, ge=1, le=168, description="匯出最近N小時的數據"),
    current_user: Dict = Depends(get_current_user_with_export_permission)
):
    """匯出指標數據為 Excel 格式"""
    
    if not EXCEL_AVAILABLE:
        raise HTTPException(
            status_code=501,
            detail="Excel 匯出功能不可用，請安裝 openpyxl: pip install openpyxl"
        )
    
    # 獲取數據
    metrics_data = get_metrics_data(hours)
    
    if not metrics_data:
        # 如果沒有數據，返回空的 Excel
        metrics_data = [{
            "timestamp": datetime.utcnow().isoformat(),
            "stream_id": "no_data",
            "people_count": 0,
            "fps_processing": 0.0,
            "fps_source": 0.0,
            "fps_inference": 0.0,
            "latency_ms": 0.0,
            "cpu_percent": 0.0,
            "memory_mb": 0.0,
            "gpu_utilization": 0.0,
            "gpu_memory_gb": 0.0,
            "total_detections": 0,
            "queue_size": 0,
            "dropped_frames": 0,
            "session_id": ""
        }]
    
    # 創建 Excel 工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "YOLO Metrics"
    
    # 添加標題行
    headers = [
        "Timestamp", "Stream ID", "People Count", "FPS Processing",
        "FPS Source", "FPS Inference", "Latency (ms)", "CPU %",
        "Memory (MB)", "GPU Utilization %", "GPU Memory (GB)", "Total Detections",
        "Queue Size", "Dropped Frames", "Session ID"
    ]
    
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # 添加數據行
    for row_idx, data in enumerate(metrics_data, 2):
        ws.cell(row=row_idx, column=1, value=data["timestamp"])
        ws.cell(row=row_idx, column=2, value=data["stream_id"])
        ws.cell(row=row_idx, column=3, value=data["people_count"])
        ws.cell(row=row_idx, column=4, value=data["fps_processing"])
        ws.cell(row=row_idx, column=5, value=data["fps_source"])
        ws.cell(row=row_idx, column=6, value=data["fps_inference"])
        ws.cell(row=row_idx, column=7, value=data["latency_ms"])
        ws.cell(row=row_idx, column=8, value=data["cpu_percent"])
        ws.cell(row=row_idx, column=9, value=data["memory_mb"])
        ws.cell(row=row_idx, column=10, value=data["gpu_utilization"])
        ws.cell(row=row_idx, column=11, value=data["gpu_memory_gb"])
        ws.cell(row=row_idx, column=12, value=data["total_detections"])
        ws.cell(row=row_idx, column=13, value=data["queue_size"])
        ws.cell(row=row_idx, column=14, value=data["dropped_frames"])
        ws.cell(row=row_idx, column=15, value=data["session_id"])
    
    # 自動調整列寬
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # 保存到內存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"yolo_metrics_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    return StreamingResponse(
        io.BytesIO(output.read()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )

@router.get("/cameras.csv")
async def export_cameras_csv(
    current_user: Dict = Depends(get_current_user_with_export_permission)
):
    """匯出攝影機列表為 CSV 格式"""
    from src.database.camera_repository import CameraRepository
    
    camera_repo = CameraRepository()
    cameras = camera_repo.list_cameras(include_source_url=True)
    
    # 創建 CSV 內容
    output = io.StringIO()
    fieldnames = [
        "id", "slug", "name", "source_url", "protocol", "description",
        "enabled", "is_online", "owner_id", "created_at", "updated_at",
        "last_tested_at", "last_test_status", "last_seen_at"
    ]
    
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    
    for camera in cameras:
        writer.writerow({
            "id": camera.get("id", ""),
            "slug": camera.get("slug", ""),
            "name": camera.get("name", ""),
            "source_url": camera.get("source_url", ""),
            "protocol": camera.get("protocol", ""),
            "description": camera.get("description", ""),
            "enabled": camera.get("enabled", False),
            "is_online": camera.get("is_online", False),
            "owner_id": camera.get("owner_id", ""),
            "created_at": camera.get("created_at", ""),
            "updated_at": camera.get("updated_at", ""),
            "last_tested_at": camera.get("last_tested_at", ""),
            "last_test_status": camera.get("last_test_status", ""),
            "last_seen_at": camera.get("last_seen_at", "")
        })
    
    csv_content = output.getvalue()
    output.close()
    
    filename = f"yolo_cameras_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
    
    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )
